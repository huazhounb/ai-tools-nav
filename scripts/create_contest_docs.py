from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "网页设计大赛"
TEMP = Path.home() / "AppData" / "Local" / "Temp"

PLAN_SRC = TEMP / "第一届“页展芳华・科创未来” 网页设计大赛实施方案.docx"
REG_SRC = TEMP / "附件1：第一届 “页展芳华・科创未来” 网页设计大赛报名表.docx"
DESC_SRC = TEMP / "附件4：第一届 “页展芳华・科创未来” 网页设计大赛作品说明文档模板 (1).docx"

PLAN_DST = OUT / "第一届“页展芳华・科创未来”网页设计大赛实施方案.docx"
REG_DST = OUT / "附件1-网页设计大赛报名表-卓晓炜.docx"
DESC_DST = OUT / "附件4-作品说明文档-AI工具箱-卓晓炜.docx"
BRIEF_DST = OUT / "AI工具箱-300字作品说明.docx"
VIDEO_GUIDE_DST = OUT / "演示视频录制脚本与提交说明.docx"

FONT_NAME = "宋体"

SUMMARY_300 = (
    "AI工具箱是一款围绕“AI + 校园・智享未来”主题设计的中文 AI 工具导航与学习平台。"
    "作品面向学生、办公族、程序员、设计师和内容创作者，整合 AI 对话、写作、绘图、编程、PPT、论文研究等工具，"
    "并提供工具详情、教程中心、提示词库、工具对比、榜单推荐和投稿审核功能。用户可以通过首页搜索框、分类筛选和热门入口快速定位适合任务的 AI 工具；"
    "在提示词页面可直接复制模板，在教程页面可按步骤学习工具用法，在对比页面可判断不同工具的适用场景。页面采用浅色蓝紫科技风，强调清晰层级、卡片化布局和移动端适配。"
    "技术上使用 Nuxt 4、Vue 3、TypeScript、Tailwind CSS、Nuxt Content 和 Supabase，兼顾 SEO 展示、登录投稿和后台审核。"
    "作品突出实用性和可扩展性，目标是帮助校园用户从“知道 AI 工具”走向“真正用 AI 完成学习和创作任务”。"
)


def set_font(run, size=10.5, bold=False):
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.bold = bold


def normalize_doc(doc):
    for p in doc.paragraphs:
        for run in p.runs:
            set_font(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_font(run)


def set_cell_text(cell, text):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_font(run)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def fill_registration():
    copy2(REG_SRC, REG_DST)
    doc = Document(str(REG_DST))
    for p in doc.paragraphs:
        if "填报日期" in p.text:
            p.text = "填报日期：2026年5月16日                  参赛形式：☑个人参赛 □团队参赛（1-3 人）"
            for run in p.runs:
                set_font(run)

    table = doc.tables[0]
    set_cell_text(table.rows[1].cells[2], "个人名称：卓晓炜")
    set_cell_text(table.rows[2].cells[2], "□文化传承・数字新表达 ☑AI + 校园・智享未来 □绿色科技・低碳同行 □未来校园・智慧生活")
    set_cell_text(table.rows[3].cells[2], "无")
    set_cell_text(
        table.rows[4].cells[2],
        "姓名：卓晓炜   性别：□男 □女   专业：计算机应用技术\n"
        "年级班级：计应三班   学号：待补充\n"
        "团队分工：独立完成网页设计、前端开发、内容整理、说明文档与演示视频制作\n"
        "联系电话：待补充        QQ / 微信：待补充\n"
        "电子邮箱：待补充",
    )
    set_cell_text(table.rows[5].cells[2], "无")
    set_cell_text(table.rows[6].cells[2], "无")
    set_cell_text(table.rows[7].cells[2], "AI工具箱")
    set_cell_text(
        table.rows[8].cells[2],
        "本作品围绕“AI + 校园・智享未来”主题，面向学生和校园用户整理 AI 工具、教程、提示词、工具对比和榜单。"
        "作品通过搜索、分类筛选、详情展示、可复制提示词和投稿审核等功能，帮助用户快速找到适合学习、写作、办公、编程和内容创作的 AI 工具，"
        "体现 AI 技术对学习效率和校园数字化服务的赋能。",
    )
    set_cell_text(table.rows[9].cells[2], "①AI工具搜索与分类筛选 ②教程/提示词/对比/榜单内容展示 ③用户投稿与管理员审核（可选）")
    set_cell_text(table.rows[10].cells[2], "Nuxt 4 + Vue 3 + TypeScript + Tailwind CSS + Nuxt Content + Supabase；开发工具：VS Code")
    set_cell_text(
        table.rows[11].cells[2],
        "所填所有信息真实有效，参赛作品为本人原创完成，未抄袭、盗用他人作品或现成模板；素材来源合法合规，遵守大赛规则。\n"
        "签字：（个人 / 组长签字）卓晓炜\n日期：2026年5月16日",
    )
    normalize_doc(doc)
    doc.save(str(REG_DST))


def fill_description():
    copy2(DESC_SRC, DESC_DST)
    doc = Document(str(DESC_DST))
    replacements = {
        "作品名称：________________________": "作品名称：AI工具箱",
        "参赛主题（勾选所选主题）：□文化传承・数字新表达 □AI + 校园・智享未来": "参赛主题（勾选所选主题）：□文化传承・数字新表达 ☑AI + 校园・智享未来",
        "团队 / 个人名称：_______________            所属学院 / 年级班级：________________": "团队 / 个人名称：卓晓炜            所属学院 / 年级班级：信息学院 / 计应三班",
        "指导教师：____________（无则填 “无”）     提交日期：______年____月____日": "指导教师：无（无则填 “无”）     提交日期：2026年5月16日",
        "核心功能一：________________________": "核心功能一：AI 工具搜索与分类筛选",
        "核心功能二：________________________": "核心功能二：教程、提示词、对比与榜单内容展示",
        "拓展功能（可选）：________________________": "拓展功能（可选）：用户投稿、登录注册与管理员审核",
        "开发工具：________________________（如 HBuilderX、VS Code、Photoshop、剪映）": "开发工具：VS Code、Chrome、剪映（用于演示视频剪辑）",
        "核心技术栈：________________________（如 HTML5+CSS3+JavaScript、Bootstrap 框架、AIGC 相关插件 / API、全景图制作工具）": "核心技术栈：Nuxt 4、Vue 3、TypeScript、Tailwind CSS、Nuxt Content、Supabase、HTML5、CSS3、JavaScript",
        "图片素材：________________________（示例：选文化主题：开平碉楼官方文旅平台、可商用国潮素材库、团队实地原创拍摄；选绿色主题：环保公益平台可商用图片、阿里巴巴矢量图标库免费图标）": "图片素材：项目主要使用 CSS 渐变、图标组件及工具官网 favicon / 官方图标来源，仅用于工具识别展示，不暗示官方合作。",
        "文字素材：________________________（示例：选 AI 校园主题：校园官网服务介绍、教育类科普平台内容，经团队整理编辑；选未来校园主题：团队原创畅想内容、校园师生需求调研整理内容）": "文字素材：工具说明、教程、提示词、对比和榜单文案由本人围绕中文 AI 工具使用场景整理撰写。",
        "音频 / 视频 / 插件素材（可选）：________________________（示例：选文化主题：非遗讲解音频来自地方非遗保护中心，可商用；选 AI 校园主题：AIGC 简易插件来自开源平台 XXX，素材链接 XXX）": "音频 / 视频 / 插件素材（可选）：演示视频由本人录屏制作；项目使用 Nuxt、Vue、Tailwind、Supabase 等开源技术生态。",
        "其他素材（如图标、模板、框架）：________________________（示例：阿里巴巴矢量图标库可商用图标、Bootstrap 官方开源框架、全景图制作开源工具）": "其他素材（如图标、模板、框架）：Lucide 图标、simple-icons、工具官方 favicon / 官方图标 URL、Nuxt 与 Tailwind CSS 开源框架。",
    }

    filled_function_one = False
    filled_member = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if text in replacements:
            p.text = replacements[text]
        elif text.startswith("选「AI + 校园・智享未来」：说明如何将"):
            p.text = (
                "本作品选择“AI + 校园・智享未来”主题，重点解决校园用户面对大量 AI 工具时“不知道选哪个、不会上手、缺少可复制模板”的问题。"
                "网站将 AI 工具选择、教程学习、场景方案、提示词复制和工具对比整合到一个平台中，帮助学生在写论文、做 PPT、编程、资料整理和内容创作等任务中更高效地使用 AI。"
            )
        elif text.startswith("选「AI + 校园・智享未来」：采用简约科技风"):
            p.text = (
                "整体采用浅色蓝紫科技风，页面布局以搜索入口、任务场景、工具卡片和内容模块为核心，强调信息层级清晰、操作路径短和移动端可读性。"
                "首页突出搜索与分类入口，详情页突出工具结论、适合人群、优缺点和相关教程，提示词卡片强调复制操作，对比页和榜单页帮助用户快速决策。"
            )
        elif text == "针对________________________功能，使用________________________技术实现":
            if not filled_function_one:
                p.text = "针对 AI 工具搜索与分类筛选功能，使用 Vue 响应式状态、computed 计算属性和 mock 数据 / API 数据兜底实现关键词搜索、分类过滤、免费状态筛选和工具卡片展示。"
                filled_function_one = True
            else:
                p.text = "针对投稿与管理员审核功能，使用 Supabase Auth 管理登录状态，服务端 API 校验权限，投稿进入审核队列，管理员可审核并发布到正式内容表。"
        elif text.startswith("（姓名）：（年级班级），主要负责"):
            if not filled_member:
                p.text = "卓晓炜：计应三班，个人参赛，独立完成所有网页设计、前端开发、内容整理、功能调试、说明文档撰写和演示视频制作。"
                filled_member = True
            else:
                p.text = ""
        for run in p.runs:
            set_font(run)

    p = doc.add_paragraph()
    run = p.add_run("七、作品介绍（300字说明）")
    set_font(run, 12, True)
    p = doc.add_paragraph(SUMMARY_300)
    for run in p.runs:
        set_font(run)
    normalize_doc(doc)
    doc.save(str(DESC_DST))


def create_brief():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    style.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("AI工具箱作品说明")
    set_font(run, 16, True)
    doc.add_paragraph("参赛人：卓晓炜    学院班级：信息学院 / 计应三班    指导教师：无")
    doc.add_paragraph(SUMMARY_300)
    doc.add_paragraph("建议参赛主题：AI + 校园・智享未来")
    doc.add_paragraph("建议提交名称：AI工具箱")
    normalize_doc(doc)
    doc.save(str(BRIEF_DST))


def create_video_guide():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    style.font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("AI工具箱演示视频录制脚本与提交说明")
    set_font(run, 16, True)
    sections = [
        ("视频时长", "建议 2-3 分钟，横屏录制，分辨率 1080p，格式 MP4。"),
        ("录制工具", "可使用 OBS、剪映录屏、Windows Xbox Game Bar 或浏览器录屏工具。录制前关闭无关窗口和通知。"),
        (
            "推荐流程",
            "1. 10秒：展示首页和作品名称，说明主题为“AI + 校园・智享未来”。\n"
            "2. 30秒：演示搜索框，输入“PPT”或“论文”，展示工具筛选结果。\n"
            "3. 35秒：进入工具详情页，展示工具结论、免费情况、中文支持和相关教程。\n"
            "4. 35秒：进入教程中心和提示词库，演示查看教程、复制提示词。\n"
            "5. 25秒：展示工具对比、榜单和免费工具入口。\n"
            "6. 25秒：演示登录后投稿 / 管理员审核入口，可简单展示表单和后台审核列表。\n"
            "7. 10秒：总结作品价值：帮助校园用户更快找到并用好 AI 工具。",
        ),
        ("注意事项", "录制前先运行项目并确认页面无报错；浏览器缩放建议 100%；讲解声音清晰；如果不录人声，可用字幕说明每一步。"),
    ]
    for title, body in sections:
        p = doc.add_paragraph()
        run = p.add_run(title)
        set_font(run, 12, True)
        for line in body.split("\n"):
            p = doc.add_paragraph(line)
            for run in p.runs:
                set_font(run)
    normalize_doc(doc)
    doc.save(str(VIDEO_GUIDE_DST))


def main():
    OUT.mkdir(exist_ok=True)
    copy2(PLAN_SRC, PLAN_DST)
    fill_registration()
    fill_description()
    create_brief()
    create_video_guide()
    for file in [PLAN_DST, REG_DST, DESC_DST, BRIEF_DST, VIDEO_GUIDE_DST]:
        print(file)


if __name__ == "__main__":
    main()
